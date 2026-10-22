# -*- coding: utf-8 -*-
"""Teste único: o endereço informado (logradouro + numeral) bate com a coordenada?

Este script existe porque as passadas anteriores responderam perto da pergunta,
mas não exatamente ela. Três lacunas que ele fecha:

  1. O NUMERAL não entrava no veredito final. A busca reversa comparava só o
     nome da via, então "Avenida São José, 100" e "Avenida São José, 4000"
     davam o mesmo resultado — a 4 km de distância. Aqui o ponto de comparação
     é o endereço COM o número.

  2. 56 das 139 linhas que não batiam ficaram sem coordenada correta. Agora toda
     linha reprovada recebe coordenada, ou uma razão explícita de por que não
     pode receber.

  3. 48 linhas que batiam carregavam coordenada de correção, resíduo das
     passadas anteriores. Se bate, não há o que corrigir — os campos saem.

REGRA: geocodifica "logradouro, numeral, bairro, município" e mede a distância
até a coordenada lançada.

    <= 150 m                          -> BATE. Coordenada certa, nada a corrigir.
    >  150 m, evidência sustentável   -> NÃO BATE. Preenche coordenada correta.
    >  150 m, evidência frágil        -> CONFERIR. Não preenche, e diz por quê.

O que torna a evidência frágil, cada uma aprendida errando antes:
    rodovia          numeral é quilômetro; o geocodificador interpola e erra km
    nome genérico    "Rua Um" existe em vários bairros e ele escolhe um calado
    nome divergente  devolveu via de outro nome — homônima ou rua renomeada
    outro município  achou a via em cidade vizinha
    acima de 5 km    dentro do mesmo município isso é falha de busca, não de
                     lançamento; adotar seria trocar um erro por um maior

Uso:
    python bate_ou_nao_bate.py
"""

import math
import os
import re
import sys
import time
import unicodedata

import pandas as pd
import psycopg2
import requests
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from _env import exigir, opcional

ENTRADA = r"C:\dados\053-26 - enderecos x coordenadas - FINAL.xlsx"
SAIDA = r"C:\dados\053-26 - enderecos x coordenadas - BATE OU NAO.xlsx"

URL = "https://maps.googleapis.com/maps/api/geocode/json"
PAUSA = 0.06
TOLERANCIA_M = 150
MAX_PLAUSIVEL_M = 5000

RE_RODOVIA = re.compile(r'\b(RODOVIA|ROD|ACESSO|SP[\s-]?\d|DUTRA|TAMOIOS|BR[\s-]?\d)\b', re.I)
RE_GENERICO = re.compile(r'^\s*(RUA|R|AVENIDA|AV|TRAVESSA|VIA)?\s*'
                         r'(UM|DOIS|TRES|TRÊS|QUATRO|CINCO|SEIS|SETE|OITO|NOVE|DEZ|ONZE|DOZE|'
                         r'[0-9]{1,2})\s*$', re.I)


def hav(a, b):
    (la1, lo1), (la2, lo2) = a, b
    R = 6371000.0
    p1, p2 = math.radians(la1), math.radians(la2)
    dp, dl = math.radians(la2 - la1), math.radians(lo2 - lo1)
    h = math.sin(dp / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * R * math.asin(math.sqrt(h))


def _tok(s):
    s = unicodedata.normalize('NFKD', str(s)).encode('ascii', 'ignore').decode().upper()
    s = re.sub(r'\b(RUA|R|AVENIDA|AV|ESTRADA|ESTR|EST|RODOVIA|ROD|TRAVESSA|TV|PRACA|PCA|'
               r'ALAMEDA|AL|VIADUTO|ACESSO|DE|DA|DO|DOS|DAS|E)\b', ' ', s)
    return {t for t in re.sub(r'[^A-Z0-9 ]', ' ', s).split() if len(t) > 2}


def nome_confere(pedido, devolvido):
    a, b = _tok(pedido), _tok(str(devolvido).split(' - ')[0].split(',')[0])
    return True if not a else len(a & b) / len(a) >= 0.5


def geocodificar(chave, endereco, tentativas=4):
    for n in range(tentativas):
        try:
            d = requests.get(URL, params={'address': endereco, 'key': chave,
                                          'region': 'br', 'language': 'pt-BR'}, timeout=25).json()
            break
        except (requests.exceptions.ConnectionError, requests.exceptions.Timeout):
            if n == tentativas - 1:
                return None
            time.sleep(2 ** n)
    time.sleep(PAUSA)
    if d.get('status') == 'OVER_QUERY_LIMIT':
        raise SystemExit("❌ Cota do Google esgotada.")
    if d.get('status') != 'OK' or not d.get('results'):
        return None
    r = d['results'][0]
    mun = None
    for c in r.get('address_components', []):
        if 'administrative_area_level_2' in c['types'] or 'locality' in c['types']:
            mun = c['long_name']
            break
    return {'endereco': r['formatted_address'], 'lat': r['geometry']['location']['lat'],
            'lon': r['geometry']['location']['lng'],
            'precisao': r['geometry']['location_type'], 'municipio': mun}


def bairros_do_banco(df):
    conn = psycopg2.connect(host=opcional('PG_HOST', 'localhost'), port=int(opcional('PG_PORT', '5432')),
                            dbname=opcional('PG_DB', '11gb'), user=opcional('PG_USER', 'admin'),
                            password=exigir('PG_PASS'))
    cur = conn.cursor()
    dt = pd.to_datetime(df['Data'], errors='coerce', dayfirst=True)
    pares = [(int(n), d.date()) for n, d in zip(df['Ocr Num'], dt)
             if pd.notna(n) and pd.notna(d)]
    cur.execute("""
        SELECT o.ocrnum, o.ocrdat, o.bairro_inicial, t.municipio
          FROM b3.ocorrencias o
          LEFT JOIN b3.tempo_resposta t ON t.id_sdov = o.id
         WHERE (o.ocrnum, o.ocrdat) = ANY(%s)
    """, (pares,))
    m = {(r[0], r[1]): (r[2], r[3]) for r in cur.fetchall()}
    conn.close()
    return m, dt


def main():
    chave = exigir('GOOGLE_MAPS_API_KEY')
    df = pd.read_excel(ENTRADA)
    banco, dt = bairros_do_banco(df)
    print(f"linhas: {len(df)}  |  casadas no banco: {len(banco)}")

    res = {c: [None] * len(df) for c in ('bate', 'lat', 'lon', 'erro', 'motivo', 'end', 'prec')}

    for n, (i, r) in enumerate(df.iterrows(), 1):
        logr = str(r['Logradouro informado (col. H)']).strip()
        num = r['Nº / km (col. I)']
        num = '' if pd.isna(num) else str(num).split('.')[0]
        k = (int(r['Ocr Num']), dt.iloc[i].date()) if pd.notna(r['Ocr Num']) and pd.notna(dt.iloc[i]) else None
        bairro, mun_bd = banco.get(k, (None, None))
        alvo_mun = mun_bd or 'São José dos Campos'

        partes = [logr] + ([num] if num else []) + ([bairro] if bairro else []) + [alvo_mun, 'SP', 'Brasil']
        g = geocodificar(chave, ', '.join(partes))

        if not g:
            res['bate'][i] = 'CONFERIR'
            res['motivo'][i] = 'Endereço não localizado na busca. Depende de consulta à guarnição.'
            continue

        res['end'][i] = g['endereco']
        res['prec'][i] = g['precisao']
        d = round(hav((g['lat'], g['lon']),
                      (float(r['Latitude informada (col. J)']), float(r['Longitude informada (col. K)']))))

        if d <= TOLERANCIA_M:
            res['bate'][i] = 'BATE'
            res['motivo'][i] = (f"Endereço e coordenada coincidem ({d} m). "
                                f"Localizado em: {g['endereco']}.")
        elif g['municipio'] and alvo_mun.lower() not in g['municipio'].lower():
            res['bate'][i] = 'CONFERIR'
            res['motivo'][i] = (f"O endereço foi localizado em {g['municipio']}, não em {alvo_mun}. "
                                f"Há via homônima na região — confirmar o município da ocorrência.")
        elif RE_RODOVIA.search(logr):
            res['bate'][i] = 'CONFERIR'
            res['motivo'][i] = (f"Rodovia ou acesso: o numeral é referência de quilômetro, não "
                                f"endereço, e a busca erra por quilômetros. Conferir pelo marco "
                                f"quilométrico ou pela referência do despacho.")
        elif RE_GENERICO.match(logr):
            res['bate'][i] = 'CONFERIR'
            res['motivo'][i] = (f"Nome genérico de via, que se repete em vários bairros do "
                                f"município. A busca escolhe uma delas sem sinalizar. Informar o "
                                f"bairro para desambiguar.")
        elif not nome_confere(logr, g['endereco']):
            res['bate'][i] = 'CONFERIR'
            res['motivo'][i] = (f"A busca devolveu '{g['endereco']}', de nome diferente do "
                                f"informado. Pode ser via renomeada, homônima ou erro de grafia "
                                f"no despacho.")
        elif d > MAX_PLAUSIVEL_M:
            res['bate'][i] = 'CONFERIR'
            res['motivo'][i] = (f"A busca devolveu ponto a {d/1000:.1f} km da coordenada lançada, "
                                f"dentro do mesmo município. Distância grande demais para ser erro "
                                f"de despacho — o endereço provavelmente não existe na base.")
        else:
            res['bate'][i] = 'NÃO BATE'
            res['lat'][i] = round(g['lat'], 6)
            res['lon'][i] = round(g['lon'], 6)
            res['erro'][i] = d
            res['motivo'][i] = (f"O endereço informado fica a {d} m da coordenada lançada. "
                                f"Localizado em: {g['endereco']} (precisão {g['precisao']}).")

        if n % 100 == 0:
            print(f"  {n} conferidas...", flush=True)

    saida = df[['Linha na planilha original', 'Município', 'Data', 'Ocr Num',
                'Referência do endereço (col. G)', 'Logradouro informado (col. H)',
                'Nº / km (col. I)', 'Latitude informada (col. J)',
                'Longitude informada (col. K)']].copy()
    saida['BATE?'] = res['bate']
    saida['LATITUDE CORRETA'] = res['lat']
    saida['LONGITUDE CORRETA'] = res['lon']
    saida['ERRO (m)'] = res['erro']
    saida['MOTIVO'] = res['motivo']
    saida['Endereço localizado'] = res['end']
    saida['Precisão da localização'] = res['prec']

    print("\n--- resultado ---")
    print(saida['BATE?'].value_counts().to_string())
    print(f"\nNÃO BATE com coordenada correta preenchida: "
          f"{saida[saida['BATE?']=='NÃO BATE']['LATITUDE CORRETA'].notna().sum()} de "
          f"{(saida['BATE?']=='NÃO BATE').sum()}")
    print(f"BATE carregando coordenada (não deve haver): "
          f"{saida[saida['BATE?']=='BATE']['LATITUDE CORRETA'].notna().sum()}")

    saida.to_excel(SAIDA, index=False, sheet_name='Endereços x coordenadas')

    wb = load_workbook(SAIDA); ws = wb.active
    larg = [10, 20, 12, 11, 28, 34, 10, 15, 15, 13, 15, 15, 11, 70, 40, 16]
    cab = Font(bold=True, color='FFFFFF', size=10); fundo = PatternFill('solid', fgColor='9E1B1B')
    for i in range(1, ws.max_column + 1):
        c = ws.cell(row=1, column=i)
        c.font = cab; c.fill = fundo
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = larg[i - 1] if i <= len(larg) else 20
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    cores = {'BATE': 'E4F1EB', 'NÃO BATE': 'FBE9EA', 'CONFERIR': 'F8EEDC'}
    for l in range(2, ws.max_row + 1):
        cor = cores.get(ws.cell(row=l, column=10).value)
        if not cor:
            continue
        f = PatternFill('solid', fgColor=cor)
        for c in range(1, ws.max_column + 1):
            ws.cell(row=l, column=c).fill = f
        ws.cell(row=l, column=14).alignment = Alignment(wrap_text=True, vertical='top')
    wb.save(SAIDA)
    print(f"\n>> gravado: {SAIDA}")


if __name__ == "__main__":
    main()
