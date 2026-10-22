# -*- coding: utf-8 -*-
"""Reduz a apuração final ao formato da planilha original.

Mesmo princípio da primeira simplificação: o arquivo de trabalho acumulou 20
colunas de rastro, que servem para auditar e atrapalham quem vai corrigir o
lançamento. Aqui ficam as nove colunas originais e mais seis.

Uma decisão que se mantém: a confiança da coordenada corrigida vai em coluna
própria. Uma coordenada obtida no número da porta e uma obtida pela referência
do despacho não valem a mesma coisa, e fundir as duas faria a planilha parecer
mais precisa do que é.

Uso:
    python planilha_simples_v2.py
"""

import os
import sys

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ENTRADA = r"C:\dados\053-26 - enderecos x coordenadas - APURADO.xlsx"
SAIDA = r"C:\dados\053-26 - enderecos x coordenadas - RESULTADO.xlsx"

VERMELHO = "9E1B1B"
CORES = {'BATE': 'E4F1EB', 'BATE (via)': 'EAF4EE', 'NÃO BATE': 'FBE9EA', 'SEM RESPOSTA': 'F8EEDC'}

COLUNAS = [
    ('Linha na planilha original', 10), ('Município', 19), ('Data', 12), ('Ocr Num', 11),
    ('Referência do endereço (col. G)', 30), ('Logradouro informado (col. H)', 34),
    ('Nº / km (col. I)', 10), ('Latitude informada (col. J)', 14),
    ('Longitude informada (col. K)', 14), ('BATE?', 14),
    ('LATITUDE CORRETA', 15), ('LONGITUDE CORRETA', 15), ('ERRO (m)', 11),
    ('O QUE ESTÁ ERRADO', 66), ('CONFIANÇA DA CORREÇÃO', 24),
]


def main():
    d = pd.read_excel(ENTRADA)

    # Acima: recalculado aqui em vez de lido do texto do motivo — depender de
    # frase para saber um fato é o tipo de acoplamento que quebra em silêncio.
    par = (d['Latitude informada (col. J)'].round(6).astype(str) + ',' +
           d['Longitude informada (col. K)'].round(6).astype(str))
    reuso = par.map(d.assign(p=par).groupby('p')['Logradouro informado (col. H)'].nunique())

    bate, oque, conf = [], [], []
    for i, r in d.iterrows():
        b = r['BATE?']
        c = r['Confiança da coordenada corrigida']
        tem = pd.notna(r['LATITUDE CORRETA'])
        n = reuso.iloc[i]

        if b == 'BATE':
            bate.append('BATE')
            oque.append(None)
            conf.append(None)
            continue

        if b == 'BATE (via, não o numeral)':
            bate.append('BATE (via)')
            oque.append('A coordenada está na via informada. O numeral não pôde ser conferido '
                        'porque o endereço não foi localizado na busca.')
            conf.append(None)
            continue

        if b == 'SEM RESPOSTA' and not tem:
            bate.append('SEM RESPOSTA')
            oque.append('Não foi possível localizar o endereço por nenhuma via — nem pelo número, '
                        'nem pelo bairro, nem pela referência do despacho.')
            conf.append(None)
            continue

        bate.append('NÃO BATE')
        if n >= 2:
            base = (f'Coordenada usada por {int(n)} logradouros diferentes da planilha — ponto '
                    f'padrão copiado de outra ocorrência, não a localização deste endereço.')
        else:
            base = 'A coordenada lançada não corresponde ao endereço informado.'

        if not tem:
            oque.append(base + ' Não foi possível localizar o endereço para propor a correção.')
            conf.append(None)
        elif c == 'porta':
            oque.append(base + ' Endereço localizado no número exato — coordenada corrigida ao lado.')
            conf.append('Porta (número exato)')
        elif isinstance(c, str) and c.startswith('entorno'):
            oque.append(base + ' Coordenada obtida pelo ponto de referência do despacho: localiza '
                               'a região, não o número.')
            conf.append('Entorno (referência do despacho)')
        elif isinstance(c, str) and c.startswith('trecho'):
            oque.append(base + ' Coordenada obtida pelo quilômetro informado na referência.')
            conf.append('Trecho de rodovia (km)')
        elif isinstance(c, str) and c.startswith('eixo'):
            oque.append(base + ' Coordenada obtida no eixo da via: localiza a rua, não o número.')
            conf.append('Eixo da via')
        else:
            # Acima: estas vieram da busca por endereço com numeral, antes de
            # existir a coluna de confiança. A precisão do próprio retorno diz
            # o que elas são, e "Endereço localizado" não dizia nada.
            p = r.get('Precisão da localização')
            if p == 'ROOFTOP':
                oque.append(base + ' Endereço localizado no número exato — coordenada corrigida ao lado.')
                conf.append('Porta (número exato)')
            elif p == 'RANGE_INTERPOLATED':
                oque.append(base + ' Número estimado dentro da quadra — confirmar antes de adotar.')
                conf.append('Quadra (número estimado)')
            else:
                oque.append(base + ' Coordenada obtida no eixo da via: localiza a rua, não o número.')
                conf.append('Eixo da via')

    s = d[[c for c, _ in COLUNAS[:9]]].copy()
    s['BATE?'] = bate
    s['LATITUDE CORRETA'] = d['LATITUDE CORRETA']
    s['LONGITUDE CORRETA'] = d['LONGITUDE CORRETA']
    s['ERRO (m)'] = d['ERRO (m)']
    s['O QUE ESTÁ ERRADO'] = oque
    s['CONFIANÇA DA CORREÇÃO'] = conf

    print(s['BATE?'].value_counts().to_string())
    nb = s[s['BATE?'].isin(['NÃO BATE', 'SEM RESPOSTA'])]
    print(f"\nnão batem: {len(nb)} | com coordenada: {nb['LATITUDE CORRETA'].notna().sum()} "
          f"| sem: {nb['LATITUDE CORRETA'].isna().sum()}")
    print(f"BATE carregando coordenada: "
          f"{s[s['BATE?'].str.startswith('BATE')]['LATITUDE CORRETA'].notna().sum()}")
    print("\n--- confiança ---")
    print(s['CONFIANÇA DA CORREÇÃO'].value_counts(dropna=False).to_string())

    s.to_excel(SAIDA, index=False, sheet_name='Endereços x coordenadas')

    wb = load_workbook(SAIDA); ws = wb.active
    cab = Font(bold=True, color='FFFFFF', size=10); fundo = PatternFill('solid', fgColor=VERMELHO)
    for i, (nome, larg) in enumerate(COLUNAS, start=1):
        c = ws.cell(row=1, column=i)
        c.font = cab; c.fill = fundo
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = larg
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    for l in range(2, ws.max_row + 1):
        cor = CORES.get(ws.cell(row=l, column=10).value)
        if not cor:
            continue
        f = PatternFill('solid', fgColor=cor)
        for c in range(1, len(COLUNAS) + 1):
            ws.cell(row=l, column=c).fill = f
        ws.cell(row=l, column=14).alignment = Alignment(wrap_text=True, vertical='top')
    wb.save(SAIDA)
    print(f"\n>> gravado: {SAIDA}")


if __name__ == "__main__":
    main()
