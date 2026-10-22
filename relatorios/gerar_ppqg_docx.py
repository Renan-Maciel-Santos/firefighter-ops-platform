# ==========================================
# GERADOR DO RELATÓRIO MODELO PPQG / IPEG — 11º GB
# ==========================================
# Monta o .docx no formato do instrumento do PPQG (mesma estrutura do
# relatório do 3º BPRv usado como referência): Perfil da Organização,
# Critérios 1-7 (descritivos) e Critério 8 (Resultados, em tabelas).
#
# Campos que dependem do 11º GB ficam marcados no documento:
#   [PREENCHER] — não existe em lugar nenhum, precisa ser escrito/levantado
#   [EDITAR]    — há rascunho ou dado sugerido, precisa revisão do Comando
#   [VALIDAR]   — número já calculado, mas o critério de cálculo precisa ser
#                 conferido contra a série histórica 2022-2025
# ==========================================

# Acima: Geração do documento Word
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# Acima: Utilitários de path
import os, sys

# Acima: Série histórica 2022-2025 já consolidada no projeto
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from dados_historico import HISTORICO, CATALOGO


# ==========================================
# CONFIGURAÇÕES
# ==========================================

SAIDA = os.path.join(os.path.dirname(os.path.abspath(__file__)), "PPQG_11GB_MODELO.docx")

# Acima: Data de corte dos números parciais de 2026 (hoje)
DATA_CORTE = "05/08/2026"

# Acima: Cores dos marcadores
COR_PREENCHER = RGBColor(0xC0, 0x00, 0x00)   # vermelho — falta tudo
COR_EDITAR    = RGBColor(0xB4, 0x6A, 0x00)   # laranja  — revisar rascunho
COR_VALIDAR   = RGBColor(0x1F, 0x4E, 0x79)   # azul     — conferir cálculo

AZUL_CAB = "1F4E79"
CINZA_SEC = "D9E2F3"


# ==========================================
# VALORES DE 2026 (parciais, extraídos do Postgres em 05/08/2026)
# ==========================================
# Acima: Só entram aqui os indicadores cuja apuração é contagem direta e
# inequívoca. Percentuais cujo denominador/critério do histórico não está
# documentado entram como VALIDAR no corpo da tabela.

VALORES_2026 = {
    "elogios_por_bombeiro":    1.2994,   # 651 elogios / 501 ativos
    "idade_media_veicular":    10.80,    # b4.vw_ppmq_idade_media
    "viaturas_baixadas_pct":   0.2907,   # b4.vw_ppmq_taxa_baixadas (50/172)
    "reclamacoes_denuncias":   7,        # sjd.vw_kpis_ano
    "acidentes_folga_bombeiro": 1,       # sjd.vw_kpis_ano
    "acidentes_viatura_vitima": 1,       # sjd.vw_kpis_ano
    "efetivo_operacional_pct": 0.804,    # b1.vw_efetivo_operacional (403/501)
}

# Acima: Indicadores cujo 2026 saiu do banco mas com fórmula a conferir
VALIDAR_2026 = {"efetivo_operacional_pct", "viaturas_baixadas_pct"}


# ==========================================
# HELPERS DE FORMATAÇÃO
# ==========================================

def fmt(valor, tipo):
    # Acima: Formata número no padrão pt-BR conforme o tipo do catálogo
    if valor is None:
        return "—"
    if tipo == "pct":
        return f"{valor * 100:.1f}".replace(".", ",") + "%"
    if tipo == "int":
        return f"{int(valor):,}".replace(",", ".")
    if tipo == "dec":
        return f"{valor:.2f}".replace(".", ",")
    if tipo == "brl":
        s = f"{valor:,.2f}".replace(",", "@").replace(".", ",").replace("@", ".")
        return "R$ " + s
    return str(valor)


def shade(cell, cor_hex):
    # Acima: Preenchimento de fundo da célula (python-docx não expõe isso direto)
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), cor_hex)
    tc_pr.append(shd)


def run_marcado(par, texto, cor, negrito=True):
    # Acima: Insere um marcador colorido de pendência
    r = par.add_run(texto)
    r.font.color.rgb = cor
    r.bold = negrito
    return r


def p_preencher(doc, orientacao):
    # Acima: Bloco padrão "escreva aqui", com a orientação do que se espera
    par = doc.add_paragraph()
    run_marcado(par, "[PREENCHER] ", COR_PREENCHER)
    r = par.add_run(orientacao)
    r.italic = True
    r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    return par


def celula(cell, texto, negrito=False, cor=None, tam=8, centro=False):
    # Acima: Escreve o conteúdo da célula com formatação controlada
    cell.text = ""
    par = cell.paragraphs[0]
    if centro:
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = par.add_run(texto)
    r.bold = negrito
    r.font.size = Pt(tam)
    if cor is not None:
        r.font.color.rgb = cor


def tabela_cabecalho(tabela, titulos, larguras):
    # Acima: Monta a linha de cabeçalho azul com as larguras fixas
    for i, t in enumerate(titulos):
        c = tabela.rows[0].cells[i]
        celula(c, t, negrito=True, cor=RGBColor(0xFF, 0xFF, 0xFF), tam=8, centro=True)
        shade(c, AZUL_CAB)
        c.width = larguras[i]


# ==========================================
# BLOCOS DO DOCUMENTO
# ==========================================

def add_titulo_secao(doc, texto, nivel=1):
    # Acima: Cabeçalho de seção usando estilos nativos (alimenta o sumário)
    h = doc.add_heading(texto, level=nivel)
    for r in h.runs:
        r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    return h


def add_sumario(doc):
    # Acima: Campo TOC — o Word preenche ao abrir (F9 / "Atualizar índice")
    par = doc.add_paragraph()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), r'TOC \o "1-3" \h \z \u')
    run_ph = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "Clique com o botão direito e escolha 'Atualizar campo' para gerar o sumário."
    run_ph.append(t)
    fld.append(run_ph)
    par._p.append(fld)


def tabela_campos(doc, linhas, larg_rotulo=Cm(6.0), larg_valor=Cm(10.5)):
    """Tabela rótulo → valor. Cada linha é (rótulo, valor, marcador).
    marcador: None (dado pronto), 'P' (preencher), 'E' (editar)."""
    t = doc.add_table(rows=len(linhas), cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (rotulo, valor, marca) in enumerate(linhas):
        c0, c1 = t.rows[i].cells
        celula(c0, rotulo, negrito=True, tam=9)
        shade(c0, CINZA_SEC)
        c0.width = larg_rotulo
        if marca == "P":
            celula(c1, "[PREENCHER] " + valor, cor=COR_PREENCHER, tam=9)
        elif marca == "E":
            celula(c1, "[EDITAR] " + valor, cor=COR_EDITAR, tam=9)
        else:
            celula(c1, valor, tam=9)
        c1.width = larg_valor
    doc.add_paragraph()
    return t


# ==========================================
# CRITÉRIO 8 — DEFINIÇÃO DOS INDICADORES
# ==========================================
# Acima: (código, chave no dados_historico OU None, rótulo, unidade, sentido)
# chave None => indicador que o 11º GB ainda não apura

RESULTADOS = [
    ("8.1", "Resultados econômico-financeiros", [
        ("8.1.1", None,                       "Índice de consumo de energia elétrica da OPM", "kWh/efetivo", "↓"),
        ("8.1.2", "uge_valor_custeio",        "Recursos de custeio recebidos", "R$", "↑"),
        ("8.1.3", "uge_valor_investimento",   "Recursos de investimento recebidos", "R$", "↑"),
        ("8.1.4", None,                       "Valor absoluto de recebimento de todas as fontes", "R$", "↑"),
        ("8.1.5", "uge_processos",            "Quantidade de processos da UGE", "qtd", "↑"),
        ("8.1.6", "viaturas_baixadas_pct",    "Taxa de viaturas baixadas", "%", "↓"),
    ]),
    ("8.2", "Resultados ambientais", [
        ("8.2.1", None,                       "Taxa de descarte de baterias em conformidade ambiental", "%", "↑"),
        ("8.2.2", None,                       "Taxa de descarte de pneus em conformidade ambiental", "%", "↑"),
        ("8.2.3", None,                       "Taxa de descarte de óleo em conformidade ambiental", "%", "↑"),
        ("8.2.4", "idade_media_veicular",     "Índice de tempo de uso da viatura (idade média)", "anos", "↓"),
    ]),
    ("8.3", "Resultados sociais", [
        ("8.3.1", "turmas_bombeiro_escola",   "Bombeiro na Escola — alunos formados", "qtd", "↑"),
        ("8.3.2", None,                       "Quantidade de ações de comunicação institucional", "qtd", "↑"),
        ("8.3.3", None,                       "Quantidade de ações sociais da OPM", "qtd", "↑"),
    ]),
    ("8.4", "Resultados relativos aos clientes e ao mercado", [
        ("8.4.1", "reclamacoes_denuncias",    "Índice de reclamações e denúncias recebidas", "qtd", "↓"),
    ]),
    ("8.5", "Resultados relativos às pessoas", [
        ("8.5.1", "lmp_outorgadas",           "Outorga de Láurea de Mérito Pessoal (LMP)", "qtd", "↑"),
        ("8.5.2", "elogios_por_bombeiro",     "Índice de elogios publicados por bombeiro", "índice", "↑"),
        ("8.5.3", None,                       "Taxa de militares que concluíram cursos de formação/aperfeiçoamento/especialização", "%", "↑"),
        ("8.5.4", "pms_punidos_pd",           "Militares punidos por Procedimento Disciplinar (PD)", "qtd", "↓"),
        ("8.5.5", "absenteismo_lts_pct",      "Taxa de absenteísmo na OPM (LTS)", "%", "↓"),
        ("8.5.6", None,                       "Quantidade de ações preventivas sobre saúde mental", "qtd", "↑"),
        ("8.5.7", None,                       "Convênios e parcerias firmadas pela OPM", "qtd", "↑"),
        ("8.5.8", "acidentes_viatura_vitima", "Acidentes de trânsito com vítima militar envolvendo viatura", "qtd", "↓"),
        ("8.5.9", "acidentes_folga_bombeiro", "Acidentes de trânsito com militar na folga", "qtd", "↓"),
    ]),
    ("8.6", "Resultados relativos a fornecimento", [
        ("8.6.1", None,                       "Taxa de entregas em desacordo com as especificações", "%", "↓"),
    ]),
    ("8.7", "Resultados relativos a produtos e processos", [
        ("8.7.1", "efetivo_operacional_pct",  "Taxa de efetivo empregado na atividade operacional", "%", "↑"),
        ("8.7.2", "atendimento_20min_pct",    "Índice de atendimento de incêndio em até 20 minutos", "%", "↑"),
        ("8.7.3", "fatais_incendio",          "Vítimas fatais em ocorrências de incêndio", "qtd", "↓"),
        ("8.7.4", "fatais_transito",          "Vítimas fatais em ocorrências de acidente de trânsito", "qtd", "↓"),
        ("8.7.5", "cobertura_territorial_pct","Ocupação da Unidade na área territorial", "%", "↑"),
        ("8.7.6", "especializacao_pct",       "Índice de especialização de bombeiros", "%", "↑"),
        ("8.7.7", "fiscalizacoes",            "SAT — fiscalizações realizadas", "qtd", "↑"),
        ("8.7.8", "dias_vistoria",            "SAT — prazo médio para vistoria", "dias", "↓"),
        ("8.7.9", "dias_analise",             "SAT — prazo médio para análise", "dias", "↓"),
        ("8.7.10","dias_analise_fat",         "SAT — prazo médio para análise de FAT", "dias", "↓"),
    ]),
]


def montar_tabela_resultados(doc, itens, larguras):
    # Acima: Uma tabela por seção 8.x, com evolução 2022-2026 + meta + referencial
    titulos = ["Cód", "Indicador", "UM", "Sent", "2022", "2023", "2024", "2025",
               f"2026¹", "Meta", "RC", "Origem RC"]
    t = doc.add_table(rows=1, cols=len(titulos))
    t.style = "Table Grid"
    tabela_cabecalho(t, titulos, larguras)

    for cod, chave, rotulo, um, sent in itens:
        linha = t.add_row()
        cs = linha.cells
        for i in range(len(titulos)):
            cs[i].width = larguras[i]

        celula(cs[0], cod, tam=8, centro=True)
        celula(cs[1], rotulo, tam=8)
        celula(cs[2], um, tam=8, centro=True)
        celula(cs[3], sent, tam=8, centro=True)

        meta = CATALOGO.get(chave, {}) if chave else {}
        tipo = meta.get("fmt", "int")
        hist = HISTORICO.get(chave, {}) if chave else {}

        # Acima: Colunas 2022-2025 vêm da série já consolidada
        for j, ano in enumerate([2022, 2023, 2024, 2025], start=4):
            if chave and ano in hist:
                celula(cs[j], fmt(hist[ano], tipo), tam=8, centro=True)
            else:
                celula(cs[j], "[PREENCHER]", cor=COR_PREENCHER, tam=7, centro=True)

        # Acima: 2026 é parcial — só preenche onde a apuração é contagem direta
        if chave and chave in VALORES_2026:
            cor = COR_VALIDAR if chave in VALIDAR_2026 else None
            txt = fmt(VALORES_2026[chave], tipo)
            if chave in VALIDAR_2026:
                txt += " [VALIDAR]"
            celula(cs[8], txt, cor=cor, tam=7 if cor else 8, centro=True)
        else:
            celula(cs[8], "[PREENCHER]", cor=COR_PREENCHER, tam=7, centro=True)

        # Acima: Meta e referencial comparativo não existem no projeto hoje
        celula(cs[9],  "[PREENCHER]", cor=COR_PREENCHER, tam=7, centro=True)
        celula(cs[10], "[PREENCHER]", cor=COR_PREENCHER, tam=7, centro=True)
        celula(cs[11], "[PREENCHER]", cor=COR_PREENCHER, tam=7, centro=True)

    doc.add_paragraph()


# ==========================================
# CRITÉRIOS 1-7 — ROTEIRO DOS SUBITENS
# ==========================================

CRITERIOS = [
    ("1. Liderança", [
        ("1.1. Desenvolvimento da Cultura", [
            ("a) Estabelecimento de valores, princípios e padrões de conduta",
             "Como o 11º GB dissemina valores e conduta: preleções, instruções continuadas, "
             "solenidades, canais internos, gestão por indicadores."),
            ("b) Promoção de mudanças culturais",
             "Que mudanças de cultura foram induzidas e como (digitalização, SEI, painéis de "
             "indicadores, orientação a dados)."),
        ]),
        ("1.2. Governança", [
            ("a) Estruturação da governança e continuidade",
             "Estrutura de comando, comissões, delegações formais e plano de continuidade."),
            ("b) Controle de riscos e conformidade",
             "Mecanismos de controle interno, auditorias, correição, tratamento de não conformidades."),
        ]),
        ("1.3. Exercício da Liderança", [
            ("a) Interação com as partes interessadas",
             "Como o Comando interage com tropa, população, órgãos parceiros e imprensa."),
            ("b) Tomada de decisão",
             "Fóruns de decisão, reuniões de análise crítica, uso de indicadores na decisão."),
        ]),
    ]),
    ("2. Estratégias", [
        ("2.1. Estratégias sustentáveis", [
            ("a) Definição da esfera de influência e objetivos", "Recorte de atuação e objetivos de longo prazo."),
            ("b) Mapeamento e tratamento de forças externas", "Análise do ambiente externo (ameaças e oportunidades)."),
            ("c) Compilação e tratamento das forças internas", "Análise interna (forças e fraquezas)."),
            ("d) Consenso estratégico", "Como as estratégias são pactuadas com a liderança."),
        ]),
        ("2.2. Implementação das estratégias", [
            ("a) Definição de mudanças", "Que mudanças foram planejadas a partir da estratégia."),
            ("b) Desdobramento dos planos", "Como as metas descem até SGB, estações e seções."),
        ]),
        ("2.3. Análise de desempenho", [
            ("a) Medição e avaliação de resultados",
             "Sistemática de medição. Citar o painel de indicadores e os ETLs automatizados do 11º GB."),
            ("b) Avaliação de progresso", "Como o avanço em relação às metas é avaliado e corrigido."),
        ]),
    ]),
    ("3. Clientes", [
        ("3.1. Mercado e atuação sustentável", [
            ("a) Estudo de mercado", "Estudo da área de atuação e da demanda por serviços de bombeiro."),
            ("b) Segmentação de mercado", "Segmentos atendidos (população, empresas, órgãos públicos)."),
            ("c) Definição de clientes-alvo", "Definição dos públicos prioritários."),
            ("d) Conhecimento sobre os clientes-alvo", "Como o 11º GB conhece necessidades e expectativas."),
            ("e) Planejamento de experiências sustentáveis", "Planejamento do atendimento ao cidadão."),
            ("f) Divulgação de produtos e características sustentáveis", "Comunicação social e campanhas."),
            ("g) Desenvolvimento de marcas sustentáveis", "Fortalecimento da imagem institucional."),
        ]),
        ("3.2. Experiência sustentável", [
            ("a) Responsividade aos clientes", "Canais e tempo de resposta ao cidadão."),
            ("b) Resolutividade de manifestações", "Tratamento de reclamações e denúncias."),
            ("c) Avaliação da satisfação dos clientes",
             "Pesquisa de satisfação. CONFERIR se este subitem existe no instrumento do ciclo vigente."),
            ("d) Fidelização dos clientes", "Ações de aproximação e confiança da população."),
        ]),
    ]),
    ("4. Sociedade", [
        ("4.1. Responsabilidade socioambiental", [
            ("a) Manutenção da legalidade", "Conformidade legal e tratamento de exigências."),
            ("b) Atuação para o desenvolvimento sustentável", "Ações socioambientais da Unidade."),
        ]),
        ("4.2. Segurança socioambiental", [
            ("a) Controle de impactos adversos", "Controle de resíduos, efluentes e impactos das operações."),
            ("b) Prontidão para emergências", "Prontidão da Unidade — este é o núcleo da atividade do 11º GB."),
        ]),
    ]),
    ("5. Conhecimento, Inovação e Tecnologia", [
        ("5.1. Conhecimento essencial", [
            ("a) Mapeamento dos conhecimentos", "Conhecimentos críticos da atividade de bombeiro."),
            ("b) Internalização do conhecimento", "Instruções, estágios, doutrina, retenção de conhecimento."),
        ]),
        ("5.2. Inovação sustentável", [
            ("a) Fomento à inovação", "Como ideias novas são captadas e testadas."),
            ("b) Experimentação em larga escala", "Pilotos e sua expansão."),
        ]),
        ("5.3. Adaptação digital", [
            ("a) Desenvolvimento digital do negócio",
             "Forte aqui: painel web do 11º GB, ETLs automatizados do SGP, mapa de força, B.I.A. no WhatsApp."),
            ("b) Segurança digital", "Controle de acesso, perfis, backup e proteção de dados pessoais."),
        ]),
    ]),
    ("6. Pessoas", [
        ("6.1. Equipes de alto desempenho", [
            ("a) Estruturação das equipes", "Organização em SGB, estações e guarnições."),
            ("b) Composição das equipes", "Perfil e dimensionamento do efetivo."),
            ("c) Desenvolvimento de competências das pessoas", "Capacitação e especialização."),
            ("d) Tratamento dos perigos e riscos de saúde e segurança",
             "Saúde ocupacional. Base disponível: IAS, TAF, restrições e afastamentos."),
            ("e) Otimização do clima organizacional", "Pesquisa de clima e ações de valorização."),
        ]),
        ("6.2. Desenvolvimento de lideranças", [
            ("a) Definição de competências de liderança", "Competências esperadas dos líderes."),
            ("b) Seleção de líderes e sucessores", "Critérios de seleção e sucessão."),
            ("c) Desenvolvimento de competências de líderes", "Formação continuada de líderes."),
        ]),
    ]),
    ("7. Processos", [
        ("7.1. Processos primários e de suporte", [
            ("a) Planejamento dos fatores de desempenho", "Definição dos fatores críticos e seus indicadores."),
            ("b) Projeto de produtos e processos sustentáveis", "Desenho dos processos finalísticos."),
            ("c) Monitoramento do desempenho", "Como cada processo é monitorado."),
            ("d) Análise e melhoria de processos e produtos", "Ciclo de melhoria contínua."),
        ]),
        ("7.2. Processos de fornecimento", [
            ("a) Desenvolvimento sustentável da cadeia de suprimentos", "Relação com fornecedores internos e externos."),
            ("b) Qualificação e seleção de fornecedores sustentáveis", "Critérios de qualificação."),
            ("c) Monitoramento do fornecimento", "Acompanhamento de entregas e prazos."),
            ("d) Avaliação do fornecimento", "Avaliação formal dos fornecedores."),
        ]),
        ("7.3. Processos econômico-financeiros", [
            ("a) Síntese de fatores de desempenho econômico-financeiro", "Indicadores financeiros acompanhados."),
            ("b) Projeção orçamentária", "Como o orçamento é projetado."),
            ("c) Qualificação de investimentos", "Critérios de priorização de investimentos."),
            ("d) Equilíbrio do fluxo financeiro", "Gestão da execução ao longo do exercício."),
        ]),
    ]),
]


# ==========================================
# PENDÊNCIAS CONSOLIDADAS
# ==========================================

PENDENCIAS = [
    ("Bloqueante", "Instrumento oficial do ciclo vigente",
     "Este modelo foi derivado do relatório do 3º BPRv referente ao PPQG 2025. É preciso obter do "
     "IPEG o instrumento do ciclo atual (critérios, limite de páginas, formulário de candidatura) "
     "e reconferir a numeração dos itens antes de escrever o conteúdo definitivo."),
    ("Bloqueante", "Referencial comparativo (RC)",
     "Cada indicador precisa ser comparado com outra unidade. Definir com qual GB o 11º se compara "
     "e obter a série de dados dessa unidade — não existe nada de outra OPM no projeto."),
    ("Bloqueante", "Metas / compromissos",
     "As colunas de meta de todos os indicadores dependem do Plano de Comando do 11º GB, que não "
     "está no projeto."),
    ("Dados", "Consumo de energia elétrica (8.1.1)",
     "kWh por efetivo, série anual. Levantar com a seção administrativa / contas de energia."),
    ("Dados", "Descarte ambiental de baterias, pneus e óleo (8.2.1 a 8.2.3)",
     "Percentual de descarte em conformidade. Levantar com MOTOMEC."),
    ("Dados", "Ações de comunicação institucional e ações sociais (8.3.2 e 8.3.3)",
     "Contagem anual. Levantar com a B/5 e a Comunicação Social."),
    ("Dados", "Cursos concluídos (8.5.3)",
     "Percentual do efetivo que concluiu cursos de formação, aperfeiçoamento, habilitação ou "
     "especialização. Levantar com a seção de ensino."),
    ("Dados", "Ações preventivas de saúde mental (8.5.6)",
     "Quantidade de ações realizadas. Atenção: o banco tem as supervisões CAPS/NAPS em "
     "b1.restricoes, mas isso é efetivo em acompanhamento, não ação realizada — são coisas "
     "diferentes e não podem ser usadas uma pela outra."),
    ("Dados", "Convênios e parcerias (8.5.7)",
     "Quantidade firmada por ano. Levantar com o Estado-Maior."),
    ("Dados", "Taxa de entregas em desacordo (8.6.1)",
     "A tabela uge.processos tem campos de entrega e prorrogação, mas com 10 registros não "
     "sustenta série histórica. Levantar com a UGE."),
    ("Dados", "Láureas de Mérito Pessoal em 2026 (8.5.1)",
     "A série 2022-2025 existe; falta o número de 2026. Levantar com a B/5."),
    ("Dados", "Militares punidos por PD em 2026 (8.5.4)",
     "A série 2022-2025 existe. O schema sjd tem os processos de 2026, mas o critério de contagem "
     "de 'punido' precisa ser alinhado ao usado no histórico."),
    ("Dados", "Quebra por subunidade",
     "O modelo do PPQG abre cada indicador por subunidade (no 11º GB seria 1º a 4º SGB + EM). A "
     "série 2022-2025 do projeto só tem o total consolidado — os valores por SGB precisam ser "
     "recuperados, ou a candidatura assume apenas o consolidado."),
    ("Dados", "Fechamento de 2026",
     f"Os números de 2026 no documento são parciais, apurados em {DATA_CORTE}. Precisam ser "
     "refeitos com o ano fechado antes da entrega."),
    ("Texto", "Perfil da Organização — itens 1 a 14",
     "Identificação (CNPJ, endereço, data de criação, comandante, responsável pela candidatura), "
     "missão, valores, visão, clientes, fornecedores, partes interessadas, redes, legalidade, "
     "ambiente competitivo, desafios, histórico de melhorias, organograma e redes internas."),
    ("Texto", "Critérios 1 a 7",
     "Cerca de 20 páginas descrevendo as práticas de gestão. É o maior volume de trabalho e "
     "depende de entrevista com os chefes de seção. O documento traz o roteiro de cada subitem."),
    ("Texto", "Glossário",
     "Lista de siglas usadas no relatório."),
    ("Anexo", "Organograma",
     "Imagem do organograma do 11º GB para o item 13 do Perfil."),
]


# ==========================================
# MONTAGEM DO DOCUMENTO
# ==========================================

def main():
    doc = Document()

    # Acima: Fonte base do documento
    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(10)

    # Acima: python-docx nasce em Letter — documento brasileiro vai em A4
    sec0 = doc.sections[0]
    sec0.page_width, sec0.page_height = Cm(21.0), Cm(29.7)
    sec0.left_margin = sec0.right_margin = Cm(2.0)

    # ---------- CAPA ----------
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("RELATÓRIO DE GESTÃO")
    r.bold = True
    r.font.size = Pt(26)
    r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("11º GRUPAMENTO DE BOMBEIROS")
    r.bold = True
    r.font.size = Pt(20)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Prêmio Paulista de Qualidade da Gestão — PPQG")
    r.font.size = Pt(14)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Instituto Paulista de Excelência da Gestão — IPEG")
    r.font.size = Pt(12)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_marcado(p, "[EDITAR] Ciclo / ano da candidatura", COR_EDITAR)

    for _ in range(6):
        doc.add_paragraph()

    # ---------- LEGENDA ----------
    add_titulo_secao(doc, "Como usar este documento", 2)
    doc.add_paragraph(
        "Este é um modelo estruturado no formato do instrumento do PPQG. O conteúdo que o projeto "
        "do 11º GB já possui foi inserido; o que falta está marcado. Os marcadores são:"
    )
    leg = doc.add_table(rows=3, cols=2)
    leg.style = "Table Grid"
    for i, (marca, cor, desc) in enumerate([
        ("[PREENCHER]", COR_PREENCHER, "Não existe no projeto. Precisa ser escrito ou levantado com a seção responsável."),
        ("[EDITAR]", COR_EDITAR, "Há um dado ou rascunho sugerido, mas precisa de revisão e aprovação do Comando."),
        ("[VALIDAR]", COR_VALIDAR, "O número foi calculado a partir do banco, mas o critério de cálculo precisa ser conferido contra a série histórica 2022-2025."),
    ]):
        celula(leg.rows[i].cells[0], marca, negrito=True, cor=cor, tam=10)
        leg.rows[i].cells[0].width = Cm(3.5)
        celula(leg.rows[i].cells[1], desc, tam=9)
        leg.rows[i].cells[1].width = Cm(13.0)
    doc.add_paragraph()
    doc.add_paragraph(
        f"Os valores de 2026 são parciais, apurados em {DATA_CORTE} a partir do banco de dados do "
        "11º GB. A série 2022-2025 vem do relatório anual consolidado da Unidade."
    ).runs[0].italic = True

    doc.add_page_break()

    # ---------- SUMÁRIO ----------
    add_titulo_secao(doc, "Sumário", 1)
    add_sumario(doc)
    doc.add_page_break()

    # ---------- PERFIL ----------
    add_titulo_secao(doc, "PERFIL DA ORGANIZAÇÃO", 1)

    add_titulo_secao(doc, "1. Identificação", 2)
    tabela_campos(doc, [
        ("Nome da organização candidata – Sigla", "11º Grupamento de Bombeiros – 11º GB", None),
        ("Mês/ano de instituição", "mês/ano de criação do 11º GB", "P"),
        ("Razão social", "", "P"),
        ("CNPJ", "", "P"),
        ("Forma de atuação", "Órgãos da Administração Pública, suas Unidades Autônomas e de Apoio", "E"),
        ("Nome do responsável pela candidatura", "", "P"),
        ("Nome do principal dirigente", "Comandante do 11º GB", "P"),
        ("E-mail", "", "P"),
        ("Endereço principal", "", "P"),
    ])

    add_titulo_secao(doc, "2. Propósito", 2)
    tabela_campos(doc, [
        ("Missão / atividade-fim", "", "P"),
        ("Valores e princípios organizacionais", "", "P"),
        ("Visão / objetivo de longo prazo", "", "P"),
    ])

    add_titulo_secao(doc, "3. Operações", 2)
    doc.add_paragraph("Porte operacional", style="Intense Quote")
    par = doc.add_paragraph()
    par.add_run(
        "O 11º Grupamento de Bombeiros conta com efetivo existente de 501 militares, dos quais 403 "
        "(80,4%) empregados na atividade operacional. A Unidade organiza-se em quatro Subgrupamentos "
        "de Bombeiros (1º a 4º SGB) e Estado-Maior/COBOM, com frota de 189 viaturas, das quais 172 "
        "titulares e idade média de 10,80 anos. "
    )
    run_marcado(par, "[PREENCHER] ", COR_PREENCHER)
    r = par.add_run(
        "complementar com: área de circunscrição em km², municípios atendidos, número de estações "
        "e postos de bombeiros, e efetivo fixado (previsto) da Unidade."
    )
    r.italic = True
    r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    doc.add_paragraph("Tecnologias", style="Intense Quote")
    p_preencher(doc, "Listar os sistemas usados pela Unidade. Sugestões a partir do que já existe no "
                     "projeto: SGP (Sistema de Gestão de Pessoas), SIOPM, COPOM on-line, SIAFEM, "
                     "SIAFÍSICO, BEC, SIPL, SEI, painel web do 11º GB e assistente B.I.A. Acrescentar "
                     "os sistemas próprios do Corpo de Bombeiros (Via Fácil Bombeiros, SISCOB, etc.).")

    doc.add_paragraph("Principais processos operacionais e produtos", style="Intense Quote")
    p_preencher(doc, "Tabela de processos primários (combate a incêndio, salvamento, APH, atendimento "
                     "pré-hospitalar, atividades técnicas/SAT) e de apoio, com resumo e produtos numerados.")

    add_titulo_secao(doc, "4. Clientes", 2)
    p_preencher(doc, "Tabela: segmentos, grupos de clientes, produtos entregues e principais "
                     "necessidades/expectativas de cada grupo.")

    add_titulo_secao(doc, "5. Fornecedores", 2)
    p_preencher(doc, "Tabela: tipos de fornecedores (internos e externos), produtos fornecidos e "
                     "necessidades. Considerar MOTOMEC, Centro de Suprimento, Diretoria de Finanças, "
                     "prefeituras e FUMREBOM.")

    add_titulo_secao(doc, "6. Força de Trabalho", 2)
    tabela_campos(doc, [
        ("Quantidade de pessoas (sem terceirizados e estagiários)", "501 militares (efetivo ativo em 05/08/2026)", None),
        ("% em relação à controladora", "", "P"),
        ("Principais necessidades e expectativas", "", "P"),
        ("Grupos de pessoas e quantidade de cada grupo", "Direção / Gerencial / Supervisão / Execução — informar o quantitativo de cada", "P"),
    ])

    add_titulo_secao(doc, "7. Outras Partes Interessadas", 2)
    p_preencher(doc, "Tabela: parte interessada, instância interlocutora e necessidades/expectativas. "
                     "Considerar CBPMESP, CCB, prefeituras, Defesa Civil, SAMU, hospitais, Polícia "
                     "Militar territorial, imprensa e a própria força de trabalho.")

    add_titulo_secao(doc, "8. Redes Externas de Atuação", 2)
    p_preencher(doc, "Tabela: rede, forma de atuação e propósito.")

    add_titulo_secao(doc, "9. Legalidade", 2)
    p_preencher(doc, "Até 15 linhas sobre o arcabouço legal (art. 144 da CF, legislação estadual do "
                     "CBPMESP, normas técnicas), pendências e sanções, impactos socioambientais e "
                     "eventuais passivos.")

    add_titulo_secao(doc, "10. Ambiente Competitivo", 2)
    p_preencher(doc, "Até 15 linhas. Citar PPMQ, Plano de Comando e o próprio PPQG como indutores de "
                     "competição positiva, e listar os referenciais comparativos usados.")

    add_titulo_secao(doc, "11. Desafios e Outros Aspectos Relevantes", 2)
    p_preencher(doc, "Lista dos principais desafios da Unidade.")

    add_titulo_secao(doc, "12. Histórico de Iniciativas e Realizações de Melhoria da Gestão", 2)
    p_preencher(doc, "Tabela ano a ano (~5 anos) com os fatos relevantes: graus de PPMQ recebidos, "
                     "certificações, criação de comissões, implantação do painel de indicadores.")

    add_titulo_secao(doc, "13. Organograma", 2)
    p_preencher(doc, "Inserir a imagem do organograma do 11º GB.")

    add_titulo_secao(doc, "14. Redes Internas", 2)
    p_preencher(doc, "Tabela: grupo, áreas e partes interessadas integrantes, coordenador. Considerar "
                     "Comissão da Qualidade, Reunião de Análise Crítica e reuniões com a força de trabalho.")

    doc.add_page_break()

    # ---------- CRITÉRIOS 1-7 ----------
    add_titulo_secao(doc, "CRITÉRIOS", 1)
    doc.add_paragraph(
        "Cada subitem abaixo traz, em itálico, a orientação do que deve ser descrito. Substitua o "
        "bloco marcado pelo texto definitivo."
    )

    for nome_criterio, subitens in CRITERIOS:
        add_titulo_secao(doc, nome_criterio, 2)
        for nome_sub, alineas in subitens:
            add_titulo_secao(doc, nome_sub, 3)
            for titulo_alinea, orientacao in alineas:
                par = doc.add_paragraph()
                r = par.add_run(titulo_alinea)
                r.bold = True
                p_preencher(doc, orientacao)

    doc.add_page_break()

    # ---------- CRITÉRIO 8 (PAISAGEM) ----------
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.LANDSCAPE
    # Acima: A4 deitado — 29,7 x 21,0 cm, com 2 cm de margem = 25,7 cm úteis
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)

    add_titulo_secao(doc, "8. RESULTADOS", 1)
    doc.add_paragraph(
        "Sent = sentido de melhora do indicador (↑ maior é melhor, ↓ menor é melhor). "
        "RC = referencial comparativo. Meta = compromisso pactuado no Plano de Comando."
    )
    p = doc.add_paragraph()
    r = p.add_run(f"¹ Valores de 2026 são parciais, apurados em {DATA_CORTE}.")
    r.italic = True
    r.font.size = Pt(8)

    # Acima: Larguras fixas somando a área útil da página em paisagem (25,7 cm = 14570 dxa)
    from docx.shared import Twips
    larguras = [Twips(w) for w in
                [700, 4212, 800, 500, 830, 830, 830, 830, 830, 900, 900, 2408]]

    for cod_sec, nome_sec, itens in RESULTADOS:
        add_titulo_secao(doc, f"{cod_sec} — {nome_sec}", 2)
        montar_tabela_resultados(doc, itens, larguras)

    # ---------- PENDÊNCIAS (RETRATO) ----------
    sec2 = doc.add_section(WD_SECTION.NEW_PAGE)
    sec2.orientation = WD_ORIENT.PORTRAIT
    sec2.page_width, sec2.page_height = Cm(21.0), Cm(29.7)
    sec2.left_margin = Cm(2.0)
    sec2.right_margin = Cm(2.0)

    add_titulo_secao(doc, "ANEXO — PENDÊNCIAS CONSOLIDADAS", 1)
    doc.add_paragraph(
        "Relação do que precisa ser providenciado para que o relatório possa ser fechado. "
        "Esta página não faz parte do relatório de candidatura — remova antes da entrega."
    )

    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    larg_pend = [Cm(2.4), Cm(4.8), Cm(9.3)]
    tabela_cabecalho(t, ["Tipo", "Pendência", "Detalhe"], larg_pend)
    for tipo, titulo, detalhe in PENDENCIAS:
        linha = t.add_row().cells
        cor = COR_PREENCHER if tipo == "Bloqueante" else None
        celula(linha[0], tipo, negrito=(tipo == "Bloqueante"), cor=cor, tam=9)
        linha[0].width = larg_pend[0]
        celula(linha[1], titulo, negrito=True, tam=9)
        linha[1].width = larg_pend[1]
        celula(linha[2], detalhe, tam=9)
        linha[2].width = larg_pend[2]

    doc.add_paragraph()
    add_titulo_secao(doc, "GLOSSÁRIO", 1)
    p_preencher(doc, "Lista de siglas utilizadas no relatório, em ordem alfabética.")

    doc.save(SAIDA)
    print(f"Documento gerado: {SAIDA}")

    # Acima: Contagem rápida do que ficou marcado, para conferência
    from docx import Document as D2
    d = D2(SAIDA)
    txt = "\n".join(p.text for p in d.paragraphs)
    for tb in d.tables:
        for row in tb.rows:
            for c in row.cells:
                txt += "\n" + c.text
    print(f"  [PREENCHER]: {txt.count('[PREENCHER]')}")
    print(f"  [EDITAR]:    {txt.count('[EDITAR]')}")
    print(f"  [VALIDAR]:   {txt.count('[VALIDAR]')}")
    print(f"  parágrafos:  {len(d.paragraphs)} | tabelas: {len(d.tables)}")


if __name__ == "__main__":
    main()
