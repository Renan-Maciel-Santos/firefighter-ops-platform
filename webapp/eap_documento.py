"""
Ordem de Serviço do EAP — preenche o modelo .docx com a turma e os escalados.

O modelo é a OS de janeiro com os trechos que mudam (o amarelo do original)
trocados por campos {{...}}. O resto do texto é o padrão do ano e fica como
está: quem muda o padrão edita o modelo no Word e sobe de novo na aba EAP.

CAMPOS
    {{nome}}           trocado pelo valor de `campos["nome"]`
    {{n}}              número do item, contado em sequência pelo documento
    {{militares}}      o parágrafo vira um por militar: "1.1. Cb PM 900209-7 Fulano;"
    {{se:pracas}}      no começo do parágrafo: ele só fica se houver turma de praças
    {{se:oficiais}}    idem, para turma de Oficiais

Os campos vão num run só no modelo padrão. Se a edição no Word partir um
campo em dois runs, o parágrafo é remontado num run único — perde negrito
interno daquele parágrafo, mas não perde o texto.
"""
import copy
import io
import re

import docx

CAMPO = re.compile(r"\{\{([a-z0-9_:]+)\}\}")
CONDICIONAIS = ("se:pracas", "se:oficiais")

# Campos que o modelo precisa ter para a OS fazer sentido. Um modelo enviado
# sem eles é recusado — melhor saber no upload do que na OS de novembro.
OBRIGATORIOS = ("numero", "militares", "periodo_ead", "assinante_nome")

# Datas e endereços preenchidos saem sempre em negrito, mesmo que o modelo
# do ano (editado no Word) traga o campo sem formatação.
NEGRITO = ("data_extenso", "periodo_ead", "presencial_pracas", "local_pracas")


def campos_do_modelo(conteudo: bytes) -> set[str]:
    d = docx.Document(io.BytesIO(conteudo))
    achados = set()
    for p in _paragrafos(d):
        achados.update(CAMPO.findall(p.text))
    return achados


def _paragrafos(d):
    yield from d.paragraphs
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                yield from c.paragraphs


def _troca(p, valores: dict):
    """Substitui {{campo}} preservando a formatação de cada run."""
    def sub(texto):
        return CAMPO.sub(lambda m: str(valores.get(m.group(1), m.group(0))), texto)

    for r in p.runs:
        if "{{" in r.text:
            r.text = sub(r.text)
    # Campo partido entre runs (edição no Word): remonta num run só.
    if CAMPO.search(p.text) and p.runs:
        texto = sub(p.text)
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)
        p.runs[0].text = texto


def _destaca(p):
    """Isola cada campo de NEGRITO num run próprio, em negrito.

    O campo costuma vir no meio de um run ("no {{local_pracas}} e incluirá");
    o run é partido em antes/campo/depois, os três com a fonte original.
    """
    for r in list(p.runs):
        m = next((m for m in CAMPO.finditer(r.text) if m.group(1) in NEGRITO), None)
        while m:
            antes, campo, depois = r.text[:m.start()], m.group(0), r.text[m.end():]
            r.text = antes
            novo = copy.deepcopy(r._r)
            r._r.addnext(novo)
            run_campo = docx.text.run.Run(novo, p)
            run_campo.text = campo
            run_campo.bold = True
            resto = copy.deepcopy(r._r)
            novo.addnext(resto)
            if not antes:      # run esvaziado pelo corte (só texto: o setter limpou o resto)
                r._r.getparent().remove(r._r)
            r = docx.text.run.Run(resto, p)
            r.text = depois
            m = next((m for m in CAMPO.finditer(r.text) if m.group(1) in NEGRITO), None)
            if not m and not depois:
                resto.getparent().remove(resto)


def _remove(p):
    p._p.getparent().remove(p._p)


def gerar(modelo: bytes, campos: dict, militares: list[str],
          tem_pracas: bool, tem_oficiais: bool) -> bytes:
    """`militares` já vem formatado e na ordem ("Cb PM 900209-7 Fulano de Tal")."""
    d = docx.Document(io.BytesIO(modelo))
    manter = {"se:pracas": tem_pracas, "se:oficiais": tem_oficiais}

    item = 0
    for p in list(d.paragraphs):
        texto = p.text

        # Parágrafo condicional: some inteiro ou perde só a marca.
        cond = next((c for c in CONDICIONAIS if "{{" + c + "}}" in texto), None)
        if cond:
            if not manter[cond]:
                _remove(p)
                continue
            _troca(p, {cond: ""})
            texto = p.text

        if "{{n}}" in texto:
            item += 1
            _troca(p, {"n": item})
            texto = p.text

        if "{{militares}}" in texto:
            _expande_lista(p, item, militares)
            continue

        if "{{" in texto:
            _destaca(p)
            _troca(p, campos)

    for s in d.sections:
        for p in list(s.header.paragraphs) + list(s.footer.paragraphs):
            if "{{" in p.text:
                _troca(p, campos)

    saida = io.BytesIO()
    d.save(saida)
    return saida.getvalue()


def _expande_lista(p, item_pai: int, militares: list[str]):
    """Um parágrafo por militar, clonando a formatação do parágrafo-campo."""
    if not militares:
        _troca(p, {"militares": "(nenhum militar escalado)"})
        return
    anterior = p._p
    for i, m in enumerate(militares, 1):
        fim = "." if i == len(militares) else ";"
        novo = copy.deepcopy(p._p)
        anterior.addnext(novo)
        anterior = novo
        par = docx.text.paragraph.Paragraph(novo, p._parent)
        _troca(par, {"militares": f"{item_pai}.{i}. {m}{fim}"})
    _remove(p)


# ---------- formatação dos dados para o texto da OS ----------

MESES = ("janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho",
         "agosto", "setembro", "outubro", "novembro", "dezembro")
MES_ABREV = ("JAN", "FEV", "MAR", "ABR", "MAI", "JUN", "JUL", "AGO", "SET",
             "OUT", "NOV", "DEZ")


def data_extenso(d) -> str:
    return f"{d.day} de {MESES[d.month - 1]} de {d.year}"


def data_militar(d) -> str:
    """26OUT26"""
    return f"{d.day:02d}{MES_ABREV[d.month - 1]}{d.year % 100:02d}"


def dtg(dt) -> str:
    """Grupo data-hora: 190800NOV26"""
    return f"{dt.day:02d}{dt.hour:02d}{dt.minute:02d}{MES_ABREV[dt.month - 1]}{dt.year % 100:02d}"


# SGP exporta "1. SGT PM"; a OS escreve "1º Sgt PM".
POSTOS = {
    "TEN CEL PM": "Ten Cel PM", "MAJ PM": "Maj PM", "CAP PM": "Cap PM",
    "1. TEN PM": "1º Ten PM", "2. TEN PM": "2º Ten PM", "ASP OF PM": "Asp Of PM",
    "SUBTEN PM": "Subten PM", "1. SGT PM": "1º Sgt PM", "2. SGT PM": "2º Sgt PM",
    "3. SGT PM": "3º Sgt PM", "CB PM": "Cb PM", "SD PM": "Sd PM",
}

MINUSCULAS = {"de", "da", "das", "do", "dos", "e"}


def nome_proprio(nome: str) -> str:
    partes = (nome or "").lower().split()
    return " ".join(p if (i and p in MINUSCULAS) else p.capitalize()
                    for i, p in enumerate(partes))


def linha_militar(posto_grad: str, re: int, dv: str | None, nome: str) -> str:
    posto = POSTOS.get((posto_grad or "").upper(), posto_grad or "")
    re_txt = f"{re}-{dv}" if dv else str(re)
    return f"{posto} {re_txt} {nome_proprio(nome)}"


def destinatario(cod_opm: int | None) -> str | None:
    """Comando a quem a OS se dirige, pela OPM do militar.

    Os NATs dos SGBs (99011x900) são da SAT, não do SGB onde ficam.
    """
    if not cod_opm:
        return None
    c = str(cod_opm)
    if c.startswith("99011") and c[6:9] == "900" and c[5] in "1234":
        return "SAT"
    if c[:6] in ("990111", "990112", "990113", "990114"):
        return f"{c[5]}º SGB"
    if c.startswith("9901192") or c == "990119100":
        return "SAT"
    if c == "990119131":
        return "COBOM"
    if c.startswith("99011"):
        return "SAdm"
    return None


ORDEM_DEST = ("1º SGB", "2º SGB", "3º SGB", "4º SGB", "SAdm", "SAT", "COBOM")


def junta(itens: list[str]) -> str:
    itens = [i for i in itens if i]
    if len(itens) <= 1:
        return "".join(itens)
    return ", ".join(itens[:-1]) + " e " + itens[-1]


def texto_destinatarios(cods: list[int | None]) -> str:
    achados = {destinatario(c) for c in cods} - {None}
    return junta([d for d in ORDEM_DEST if d in achados])
